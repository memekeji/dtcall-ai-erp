from .forms import (
    ContractCategoryForm, ProductCategoryForm, ServiceCategoryForm,
    ServiceForm, SupplierForm, PurchaseCategoryForm, PurchaseItemForm, ProductForm
)
from .models import (
    ContractCategory, ProductCategory, ServiceCategory,
    Product, Service, Supplier, PurchaseCategory, PurchaseItem,
    ContractCate, ProductCate
)
from django.shortcuts import render, get_object_or_404
from django.http import Http404
from django.views import View
from rest_framework import serializers, viewsets
from django.core.paginator import Paginator
from .models import Contract, Product, Purchase
from django.http import JsonResponse, HttpResponse
from django.utils import timezone
from django.db.models import Q
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from urllib.parse import quote
import json
import logging
import time
from datetime import datetime
from decimal import Decimal, InvalidOperation

from django.contrib.auth.mixins import LoginRequiredMixin

from apps.common.utils import (
    timestamp_to_date, safe_int, parse_date_range
)
from apps.common.constants import ApiResponseCode, CommonConstant
from apps.user.models import Admin

logger = logging.getLogger(__name__)


def _format_datetime_value(value, format_str='%Y-%m-%d %H:%M'):
    if not value:
        return ''
    if isinstance(value, (int, float)):
        return timestamp_to_date(value, format_str)
    if hasattr(value, 'strftime'):
        return value.strftime(format_str)
    return str(value)


def _to_decimal(value):
    try:
        return Decimal(str(value or 0))
    except (InvalidOperation, TypeError, ValueError):
        return Decimal('0')


def _filter_active(queryset):
    model = queryset.model
    if model == Contract:
        return queryset.filter(delete_time=CommonConstant.DELETE_TIME_ZERO)
    if model == ContractCate:
        return queryset.filter(delete_time=CommonConstant.DELETE_TIME_ZERO, status=1)
    if model == ProductCate:
        return queryset.filter(delete_time__isnull=True, status=1)
    if hasattr(model, 'delete_time'):
        return queryset.filter(delete_time__isnull=True)
    if hasattr(model, 'is_active'):
        return queryset.filter(is_active=True)
    return queryset


def _get_active_object(model, object_id):
    return get_object_or_404(_filter_active(model.objects.all()), id=object_id)


def _build_choice_queryset(model, current_id=None):
    queryset = _filter_active(model.objects.all())
    if current_id:
        queryset = queryset | model.objects.filter(id=current_id)
    return queryset.distinct()


def _serialize_basic_data_response(queryset, request, formatter, order_by='-created_at', search_fields=None):
    params = request.GET.dict()
    search = params.get('search') or params.get('keywords')
    if search:
        search_fields = search_fields or ['name', 'code']
        conditions = Q()
        for field in search_fields:
            conditions |= Q(**{f'{field}__icontains': search})
        queryset = queryset.filter(conditions)

    page = safe_int(params.get('page'), 1)
    limit = safe_int(params.get('limit'), CommonConstant.DEFAULT_PAGE_SIZE)
    paginator = Paginator(queryset.order_by(order_by), limit)
    page_obj = paginator.get_page(page)

    return JsonResponse({
        'code': 0,
        'msg': '',
        'count': paginator.count,
        'data': [formatter(item) for item in page_obj]
    })


def _soft_delete_object(model, object_id):
    obj = _get_active_object(model, object_id)
    if hasattr(model, 'delete_time'):
        value = int(time.time()) if model in [Contract, ContractCate] else timezone.now()
        setattr(obj, 'delete_time', value)
        obj.save(update_fields=['delete_time'])
    elif hasattr(model, 'is_active'):
        obj.is_active = False
        obj.save(update_fields=['is_active'])
    else:
        obj.delete()
    return obj


def _json_success(message='操作成功'):
    return JsonResponse({'success': True, 'code': 0, 'message': message, 'msg': message})


def _json_error(message):
    return JsonResponse({'success': False, 'code': 1, 'message': message, 'msg': message})


def _get_basic_data_model(model_name):
    model_map = {
        'contract_category': ContractCategory,
        'product_category': ProductCategory,
        'product': Product,
        'service_category': ServiceCategory,
        'service': Service,
        'supplier': Supplier,
        'purchase_category': PurchaseCategory,
        'purchase_item': PurchaseItem,
    }
    return model_map.get(model_name)


class ContractBasicDataDeleteView(LoginRequiredMixin, View):
    login_url = '/user/login/'

    def post(self, request, model_name, id):
        model = _get_basic_data_model(model_name)
        if not model:
            return _json_error('无效的数据类型')
        try:
            _soft_delete_object(model, id)
            return _json_success('删除成功')
        except Exception as e:
            logger.error(f'删除合同基础数据失败: {str(e)}', exc_info=True)
            return _json_error('删除失败，请确认数据状态后重试')


def _parse_json_request(request):
    if not request.body:
        return {}
    try:
        return json.loads(request.body.decode('utf-8'))
    except (json.JSONDecodeError, UnicodeDecodeError):
        return {}


def _get_request_params(request):
    if request.method == 'POST':
        return request.POST
    return request.GET


def _is_data_request(request):
    return (
        request.headers.get('X-Requested-With') == 'XMLHttpRequest'
        or request.path.rstrip('/').endswith('datalist')
        or request.path.rstrip('/').endswith('archivelist')
        or request.path.rstrip('/').endswith('stoplist')
        or request.path.rstrip('/').endswith('voidlist')
    )


def _base_active_contract_queryset():
    return Contract.objects.filter(
        delete_time=CommonConstant.DELETE_TIME_ZERO,
        archive_time=CommonConstant.DELETE_TIME_ZERO,
        stop_time=CommonConstant.DELETE_TIME_ZERO,
        void_time=CommonConstant.DELETE_TIME_ZERO
    )


def _parse_int_ids(value):
    if value is None:
        return []
    if not isinstance(value, (list, tuple, set)):
        value = [value]
    ids = []
    for item in value:
        item_id = safe_int(item, 0)
        if item_id > 0 and item_id not in ids:
            ids.append(item_id)
    return ids


def _normalize_date_text(value):
    value = (value or '').strip()
    if not value:
        return None
    for format_str in ['%Y-%m-%d', '%Y-%m-%d %H:%M:%S']:
        try:
            return datetime.strptime(value, format_str).date()
        except ValueError:
            continue
    return None


def _normalize_decimal_text(value):
    decimal_value = _to_decimal(value)
    if decimal_value < 0:
        return None
    return decimal_value


def _normalize_user_id(value):
    user_id = safe_int(value, 0)
    return user_id if user_id > 0 else None


def _get_active_contract_categories():
    return ContractCate.objects.filter(
        status=1,
        delete_time=CommonConstant.DELETE_TIME_ZERO
    ).order_by('id')


def _get_purchase_form_context(purchase=None):
    return {
        'purchase': purchase,
        'contract_categories': _get_active_contract_categories(),
        'admin_users': Admin.objects.filter(status=1).order_by('id')
    }


def _parse_contract_operation_request(request, action=None):
    data = _parse_json_request(request)
    if action:
        data['action'] = action
    contract_id = safe_int(data.get('id'), 0)
    if contract_id <= 0:
        raise ValueError('缺少合同ID')
    try:
        contract = get_object_or_404(
            Contract,
            id=contract_id,
            delete_time=CommonConstant.DELETE_TIME_ZERO
        )
    except Http404:
        raise ValueError('未找到可操作的合同记录')
    return data, contract


def _build_purchase_payload(params, request_user=None, require_all=True):
    field_names = [
        'name', 'code', 'cate_id', 'types', 'amount', 'sign_time',
        'start_time', 'end_time', 'check_status', 'remark', 'share_ids',
        'check_uids', 'check_history_uids', 'file_ids'
    ]
    payload = {}

    for field in field_names:
        if field in params:
            payload[field] = params.get(field)

    for field in ['cate_id', 'types', 'check_status']:
        if field in payload:
            payload[field] = safe_int(payload.get(field), 0)

    for field in ['amount']:
        if field in payload:
            amount = _normalize_decimal_text(payload.get(field))
            if amount is None:
                raise ValueError('采购金额不能小于0')
            payload[field] = amount

    for field in ['sign_time', 'start_time', 'end_time']:
        if field in payload:
            date_value = _normalize_date_text(payload.get(field))
            if require_all and not date_value:
                raise ValueError('请选择有效的日期')
            if date_value:
                payload[field] = date_value
            else:
                payload.pop(field, None)

    for field in ['prepared_uid', 'sign_uid', 'keeper_uid']:
        if field in params:
            payload[f'{field}_id'] = _normalize_user_id(params.get(field))

    if require_all:
        required_fields = ['code', 'name', 'cate_id', 'amount', 'sign_time', 'start_time', 'end_time']
        if any(not payload.get(field) for field in required_fields):
            raise ValueError('请完整填写采购合同必填信息')
        if payload['end_time'] <= payload['start_time']:
            raise ValueError('结束时间必须大于开始时间')
    elif 'start_time' in payload and 'end_time' in payload and payload['end_time'] <= payload['start_time']:
        raise ValueError('结束时间必须大于开始时间')

    if request_user and getattr(request_user, 'is_authenticated', False):
        payload['admin_id'] = request_user.id
        for field in ['prepared_uid_id', 'sign_uid_id', 'keeper_uid_id']:
            if field not in payload or payload[field] is None:
                payload[field] = request_user.id
    elif not require_all:
        for field in ['prepared_uid_id', 'sign_uid_id', 'keeper_uid_id']:
            if field in payload and payload[field] is None:
                payload.pop(field, None)

    return payload


def _build_contract_payload(params, request_user=None, require_all=True):
    field_names = [
        'pid', 'code', 'name', 'cate_id', 'types', 'subject_id',
        'customer_id', 'customer', 'contact_name', 'contact_mobile',
        'contact_address', 'start_time', 'end_time', 'prepared_uid',
        'sign_uid', 'keeper_uid', 'share_ids', 'file_ids', 'sign_time',
        'did', 'cost', 'content', 'is_tax', 'tax', 'remark',
        'check_status', 'check_flow_id', 'check_step_sort', 'check_uids',
        'check_last_uid', 'check_history_uids', 'check_copy_uids'
    ]
    payload = {}

    for field in field_names:
        if field in params:
            payload[field] = params.get(field)

    int_fields = [
        'pid', 'cate_id', 'types', 'customer_id', 'prepared_uid',
        'sign_uid', 'keeper_uid', 'did', 'is_tax', 'check_status',
        'check_flow_id', 'check_step_sort'
    ]
    for field in int_fields:
        if field in payload:
            payload[field] = safe_int(payload.get(field), 0)

    for field in ['cost', 'tax']:
        if field in payload:
            decimal_value = _normalize_decimal_text(payload.get(field))
            if decimal_value is None:
                raise ValueError('金额不能小于0')
            payload[field] = decimal_value

    for field in ['sign_time', 'start_time', 'end_time']:
        if field in payload:
            date_value = _normalize_date_text(payload.get(field))
            if date_value:
                payload[field] = int(time.mktime(date_value.timetuple()))
            elif require_all:
                raise ValueError('请选择有效的日期')
            else:
                payload.pop(field, None)

    if payload.get('is_tax') != 1:
        payload['tax'] = Decimal('0')

    if require_all:
        required_fields = ['code', 'name', 'customer', 'cost', 'sign_time', 'start_time', 'end_time']
        if any(payload.get(field) in [None, ''] for field in required_fields):
            raise ValueError('请完整填写合同必填信息')
        if payload['end_time'] <= payload['start_time']:
            raise ValueError('结束时间必须大于开始时间')
    elif 'start_time' in payload and 'end_time' in payload and payload['end_time'] <= payload['start_time']:
        raise ValueError('结束时间必须大于开始时间')

    if request_user and getattr(request_user, 'is_authenticated', False):
        payload['admin_id'] = request_user.id

    return payload


def _build_xlsx_response(title, headers, rows, filename_prefix):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    header_fill = PatternFill('solid', fgColor='F2F2F2')
    header_font = Font(name='Arial', bold=True)
    body_font = Font(name='Arial')
    alignment = Alignment(vertical='center')
    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = alignment
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = alignment
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            value = '' if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        ws.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 36)
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f'{filename_prefix}_{datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx'
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{quote(filename)}"
    wb.save(response)
    return response


def _filter_contract_queryset(request, queryset, date_field_map=None):
    params = _get_request_params(request)
    code = params.get('code', '').strip()
    status = (params.get('status') or params.get('check_status') or '').strip()
    customer = params.get('customer', '').strip()
    keywords = params.get('keywords', '').strip()
    tab = params.get('tab', '').strip()
    uid = getattr(request.user, 'id', 0)
    tab_filter_map = {
        '1': ('admin_id', uid),
        '2': ('check_status', 0),
        '3': ('check_status__in', [2, 3]),
    }

    if code:
        queryset = queryset.filter(code__icontains=code)
    if status:
        queryset = queryset.filter(check_status=status)
    if customer:
        queryset = queryset.filter(customer__icontains=customer)
    if keywords:
        queryset = queryset.filter(Q(name__icontains=keywords) | Q(code__icontains=keywords))
    if params.get('types'):
        queryset = queryset.filter(types=params.get('types'))
    if params.get('cate_id'):
        queryset = queryset.filter(cate_id=params.get('cate_id'))

    if tab in tab_filter_map:
        field, value = tab_filter_map[tab]
        if value is not None:
            queryset = queryset.filter(**{field: value})
    elif tab == '4' and uid:
        queryset = queryset.filter(check_copy_uids__contains=str(uid))
    elif tab == '0' and request.method == 'POST' and not params.get('uid'):
        if not request.user.has_perm('contract.admin'):
            queryset = queryset.filter(
                Q(admin_id=uid) | Q(check_uids__contains=str(uid))
            )

    if date_field_map:
        for param_name, field_name in date_field_map.items():
            date_range = params.get(param_name, '').strip()
            start_time, end_time = parse_date_range(date_range)
            if start_time and end_time:
                queryset = queryset.filter(**{f'{field_name}__range': (start_time, end_time)})

    for param_name, field_name in {'sign_time': 'sign_time', 'end_time': 'end_time'}.items():
        date_range = params.get(param_name, '').strip()
        if '~' in date_range:
            start, end = date_range.split('~', 1)
            try:
                start_time = int(datetime.strptime(start.strip(), '%Y-%m-%d').timestamp())
                end_time = int(datetime.strptime(end.strip(), '%Y-%m-%d').replace(hour=23, minute=59, second=59).timestamp())
                queryset = queryset.filter(**{f'{field_name}__range': (start_time, end_time)})
            except ValueError:
                pass

    return queryset


def _build_contract_export_rows(queryset, extra_columns=None):
    headers = ['ID', '合同编号', '合同名称', '客户名称', '合同金额', '审核状态', '签订时间', '开始时间', '结束时间', '创建时间']
    extra_columns = extra_columns or []
    headers.extend([column[0] for column in extra_columns])
    status_map = {0: '待审核', 1: '审核中', 2: '审核通过', 3: '审核不通过', 4: '撤销审核'}
    rows = []
    for contract in queryset:
        row = [
            contract.id,
            contract.code or '',
            contract.name or '',
            contract.customer or '',
            float(_to_decimal(contract.cost)),
            status_map.get(contract.check_status, contract.check_status),
            timestamp_to_date(contract.sign_time),
            timestamp_to_date(contract.start_time),
            timestamp_to_date(contract.end_time),
            _format_datetime_value(contract.create_time),
        ]
        for _, value_getter in extra_columns:
            row.append(value_getter(contract))
        rows.append(row)
    return headers, rows


def _filter_purchase_queryset(request, queryset):
    params = _get_request_params(request)
    code = params.get('code', '').strip()
    status = (params.get('status') or params.get('check_status') or '').strip()
    keywords = params.get('keywords', '').strip()
    date_range = params.get('sign_date', '').strip() or params.get('date', '').strip()
    tab = params.get('tab', '').strip()
    uid = getattr(request.user, 'id', 0)
    tab_filter_map = {
        '1': ('admin_id', uid),
        '2': ('check_status', 0),
        '3': ('check_status__in', [2, 3]),
    }

    if code:
        queryset = queryset.filter(code__icontains=code)
    if status:
        queryset = queryset.filter(check_status=status)
    if keywords:
        queryset = queryset.filter(Q(name__icontains=keywords) | Q(code__icontains=keywords))
    if tab in tab_filter_map:
        field, value = tab_filter_map[tab]
        if value is not None:
            queryset = queryset.filter(**{field: value})
    if date_range and ' - ' in date_range:
        start_text, end_text = date_range.split(' - ', 1)
        try:
            start_date = datetime.strptime(start_text.strip(), '%Y-%m-%d').date()
            end_date = datetime.strptime(end_text.strip(), '%Y-%m-%d').date()
            queryset = queryset.filter(sign_time__range=(start_date, end_date))
        except ValueError:
            pass
    return queryset

try:
    FINANCE_MODULE_AVAILABLE = True
except ImportError:
    FINANCE_MODULE_AVAILABLE = False

try:
    CUSTOMER_MODULE_AVAILABLE = True
except ImportError:
    CUSTOMER_MODULE_AVAILABLE = False


# ==================== API ViewSets ====================

class ContractCategorySerializer(serializers.ModelSerializer):
    class Meta:
        model = ContractCategory
        fields = [
            'id',
            'name',
            'code',
            'parent',
            'description',
            'template_path',
            'sort_order',
            'is_active',
            'created_at']


class ContractCategoryViewSet(viewsets.ModelViewSet):
    queryset = ContractCategory.objects.filter(is_active=True)
    serializer_class = ContractCategorySerializer


class ProductCategorySerializer(serializers.ModelSerializer):
    class Meta:
        model = ProductCategory
        fields = [
            'id',
            'name',
            'code',
            'parent',
            'description',
            'sort_order',
            'is_active',
            'created_at']


class ProductCategoryViewSet(viewsets.ModelViewSet):
    queryset = ProductCategory.objects.filter(is_active=True)
    serializer_class = ProductCategorySerializer


class ServiceCategorySerializer(serializers.ModelSerializer):
    class Meta:
        model = ServiceCategory
        fields = [
            'id',
            'name',
            'code',
            'parent',
            'description',
            'sort_order',
            'is_active',
            'created_at']


class ServiceCategoryViewSet(viewsets.ModelViewSet):
    queryset = ServiceCategory.objects.filter(is_active=True)
    serializer_class = ServiceCategorySerializer


class ProductSerializer(serializers.ModelSerializer):
    category_name = serializers.CharField(source='cate.title', read_only=True)

    class Meta:
        model = Product
        fields = [
            'id',
            'code',
            'cate',
            'category_name',
            'name',
            'specs',
            'unit',
            'price',
            'remark',
            'create_time',
            'update_time',
            'delete_time']


class ProductViewSet(viewsets.ModelViewSet):
    queryset = Product.objects.filter(delete_time__isnull=True)
    serializer_class = ProductSerializer


class ServiceSerializer(serializers.ModelSerializer):
    category_name = serializers.CharField(source='category.name', read_only=True)

    class Meta:
        model = Service
        fields = [
            'id',
            'code',
            'name',
            'category',
            'category_name',
            'unit',
            'price',
            'duration',
            'description',
            'requirements',
            'is_active',
            'created_at']


class ServiceViewSet(viewsets.ModelViewSet):
    queryset = Service.objects.filter(is_active=True)
    serializer_class = ServiceSerializer


class SupplierSerializer(serializers.ModelSerializer):
    class Meta:
        model = Supplier
        fields = [
            'id',
            'name',
            'code',
            'contact_person',
            'contact_phone',
            'contact_email',
            'address',
            'tax_number',
            'bank_account',
            'bank_name',
            'credit_level',
            'business_scope',
            'is_active',
            'created_at']


class SupplierViewSet(viewsets.ModelViewSet):
    queryset = Supplier.objects.filter(is_active=True)
    serializer_class = SupplierSerializer


class PurchaseCategorySerializer(serializers.ModelSerializer):
    class Meta:
        model = PurchaseCategory
        fields = [
            'id',
            'name',
            'code',
            'parent',
            'description',
            'sort_order',
            'is_active',
            'created_at']


class PurchaseCategoryViewSet(viewsets.ModelViewSet):
    queryset = PurchaseCategory.objects.filter(is_active=True)
    serializer_class = PurchaseCategorySerializer


class PurchaseItemSerializer(serializers.ModelSerializer):
    class Meta:
        model = PurchaseItem
        fields = [
            'id',
            'name',
            'code',
            'category',
            'specification',
            'unit',
            'reference_price',
            'supplier',
            'description',
            'is_active',
            'created_at']


class PurchaseItemViewSet(viewsets.ModelViewSet):
    queryset = PurchaseItem.objects.filter(is_active=True)
    serializer_class = PurchaseItemSerializer


class ProductView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        if _is_data_request(request):
            return self.get_data_list(request)
        return render(request, 'contract/product_list.html')

    def get_data_list(self, request):
        params = _get_request_params(request)
        queryset = Product.objects.select_related('cate').filter(delete_time__isnull=True)

        if 'keywords' in params:
            queryset = queryset.filter(
                Q(name__icontains=params['keywords']) |
                Q(code__icontains=params['keywords'])
            )

        if 'cate_id' in params:
            queryset = queryset.filter(cate_id=params['cate_id'])

        page = safe_int(params.get('page'), 1)
        limit = safe_int(params.get('limit'), CommonConstant.DEFAULT_PAGE_SIZE)
        order_by = params.get('order_field', '-id')

        paginator = Paginator(queryset.order_by(order_by), limit)
        page_obj = paginator.get_page(page)

        data = []
        for product in page_obj:
            data.append({
                'id': product.id,
                'name': product.name or '',
                'code': product.code or '',
                'category': product.cate.title if product.cate else '',
                'specification': product.specs or '',
                'unit': product.unit or '',
                'price': str(product.price) if product.price else '0',
                'sort_order': 0,
                'is_active': True,
                'created_at': product.create_time.strftime('%Y-%m-%d %H:%M:%S') if product.create_time else ''
            })

        return JsonResponse({
            'code': 0,
            'msg': '',
            'count': paginator.count,
            'data': data
        })


class ProductAddView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        form = ProductForm()
        return render(request, 'contract/product_form.html', {'form': form, 'object': None})

    def post(self, request):
        form = ProductForm(request.POST)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})

        try:
            obj = form.save(commit=False)
            if hasattr(obj, 'admin_id') and getattr(obj, 'admin_id', None) in [None, 0, '']:
                obj.admin = request.user
            obj.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'添加产品失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查产品信息后重试'})


class ProductDetailView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request, id):
        product = _get_active_object(Product, id)
        form = ProductForm(instance=product)
        return render(request, 'contract/product_form.html', {'form': form, 'object': product})

    def post(self, request, id):
        product = _get_active_object(Product, id)
        form = ProductForm(request.POST, instance=product)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})

        try:
            form.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'保存产品失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查产品信息后重试'})


class ServicesView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        if _is_data_request(request):
            return self.get_data_list(request)
        return render(request, 'contract/service_list.html')

    def get_data_list(self, request):
        params = _get_request_params(request)
        queryset = Service.objects.select_related('category').filter(is_active=True)

        search = params.get('search') or params.get('keywords')
        if search:
            queryset = queryset.filter(Q(name__icontains=search) | Q(code__icontains=search))

        page = safe_int(params.get('page'), 1)
        limit = safe_int(params.get('limit'), CommonConstant.DEFAULT_PAGE_SIZE)

        paginator = Paginator(queryset.order_by('-created_at'), limit)
        page_obj = paginator.get_page(page)

        data = []
        for service in page_obj:
            data.append({
                'id': service.id,
                'name': service.name or '',
                'code': service.code or '',
                'price': str(service.price) if service.price else '0',
                'unit': service.unit or '',
                'category': service.category.name if service.category else '',
                'description': service.description or '',
                'sort_order': 0,
                'is_active': bool(service.is_active),
                'created_at': service.created_at.strftime('%Y-%m-%d %H:%M:%S') if service.created_at else ''
            })

        return JsonResponse({
            'code': 0,
            'msg': '',
            'count': paginator.count,
            'data': data
        })


class ServicesAddView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        form = ServiceForm()
        return render(request, 'contract/service_form.html', {'form': form, 'object': None})

    def post(self, request):
        form = ServiceForm(request.POST)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})

        try:
            form.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'保存服务失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查服务信息后重试'})


class ServicesDetailView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request, id):
        service = _get_active_object(Service, id)
        form = ServiceForm(instance=service)
        return render(request, 'contract/service_form.html', {'form': form, 'object': service})

    def post(self, request, id):
        service = _get_active_object(Service, id)
        form = ServiceForm(request.POST, instance=service)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})

        try:
            form.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'保存服务失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查服务信息后重试'})


class PurchaseView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        if _is_data_request(request):
            return self.get_data_list(request)
        return render(request, 'contract/purchase_list.html')

    def get_data_list(self, request):
        params = _get_request_params(request)
        queryset = _filter_purchase_queryset(
            request,
            Purchase.objects.select_related('cate').filter(delete_time__isnull=True)
        )

        page = safe_int(params.get('page'), 1)
        limit = safe_int(params.get('limit'), CommonConstant.DEFAULT_PAGE_SIZE)

        paginator = Paginator(queryset.order_by('-create_time'), limit)
        page_obj = paginator.get_page(page)

        data = []
        for purchase in page_obj:
            data.append({
                'id': purchase.id,
                'name': purchase.name or '',
                'code': purchase.code or '',
                'category': purchase.cate.title if purchase.cate else '',
                'customer': '',
                'cost': str(purchase.amount or 0),
                'amount': str(purchase.amount or 0),
                'status': purchase.check_status,
                'check_status': purchase.check_status,
                'sign_time': purchase.sign_time.strftime('%Y-%m-%d') if purchase.sign_time else '',
                'start_time': purchase.start_time.strftime('%Y-%m-%d') if purchase.start_time else '',
                'end_time': purchase.end_time.strftime('%Y-%m-%d') if purchase.end_time else '',
                'create_time': _format_datetime_value(purchase.create_time),
            })

        return JsonResponse({
            'code': ApiResponseCode.CODE_SUCCESS,
            'msg': 'success',
            'count': paginator.count,
            'data': data
        })


class PurchaseAddView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        return render(request,
                      'contract/purchase_add.html',
                      _get_purchase_form_context())

    def post(self, request):
        params = request.POST.dict()

        try:
            payload = _build_purchase_payload(params, request.user)
            Purchase.objects.create(**payload)
            logger.info(f"用户 {request.user.id} 添加了采购合同")
            return JsonResponse({
                'code': ApiResponseCode.CODE_SUCCESS,
                'msg': '添加成功'
            })
        except Exception as e:
            logger.error(f'添加采购合同失败: {str(e)}', exc_info=True)
            return JsonResponse({
                'code': ApiResponseCode.CODE_ERROR,
                'msg': '添加失败，请检查采购合同信息后重试'
            })


class PurchaseDetailView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request, id):
        purchase = get_object_or_404(Purchase, id=id, delete_time__isnull=True)
        return render(request,
                      'contract/purchase_add.html',
                      _get_purchase_form_context(purchase))


class PurchaseUpdateView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def post(self, request, id):
        try:
            purchase = get_object_or_404(Purchase, id=id, delete_time__isnull=True)
            params = request.POST.dict()
            payload = _build_purchase_payload(params, request.user, require_all=False)
            if not payload:
                return JsonResponse({
                    'code': ApiResponseCode.CODE_ERROR,
                    'msg': '没有可更新的采购合同信息'
                })
            for field, value in payload.items():
                setattr(purchase, field, value)
            purchase.save()
            logger.info(f"用户 {request.user.id} 更新了采购合同: {id}")
            return JsonResponse({
                'code': ApiResponseCode.CODE_SUCCESS,
                'msg': '更新成功'
            })
        except Exception as e:
            logger.error(f'更新采购合同失败: {str(e)}', exc_info=True)
            return JsonResponse({
                'code': ApiResponseCode.CODE_ERROR,
                'msg': '更新失败，请检查采购合同信息后重试'
            })


class PurchaseDeleteView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def post(self, request):
        try:
            data = _parse_json_request(request)
            ids = _parse_int_ids(data.get('ids') or data.get('id'))
            if not ids:
                return JsonResponse({
                    'code': ApiResponseCode.CODE_ERROR,
                    'msg': '缺少采购ID'
                })
            updated_count = Purchase.objects.filter(
                id__in=ids,
                delete_time__isnull=True
            ).update(delete_time=timezone.now())
            if updated_count == 0:
                return JsonResponse({
                    'code': ApiResponseCode.CODE_ERROR,
                    'msg': '未找到可删除的采购记录'
                })
            logger.info(f"用户 {request.user.id} 删除了采购合同: {ids}")
            return JsonResponse({
                'code': ApiResponseCode.CODE_SUCCESS,
                'msg': '删除成功'
            })
        except Exception as e:
            logger.error(f'删除采购合同失败: {str(e)}', exc_info=True)
            return JsonResponse({
                'code': ApiResponseCode.CODE_ERROR,
                'msg': '删除失败，请确认采购合同状态后重试'
            })


class ArchiveListView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        if _is_data_request(request):
            return self.get_data_list(request)
        return render(request, 'contract/archive_list.html')

    def get_data_list(self, request):
        page = safe_int(request.GET.get('page'), 1)
        limit = safe_int(
            request.GET.get('limit'),
            CommonConstant.DEFAULT_PAGE_SIZE)

        queryset = _filter_contract_queryset(
            request,
            Contract.objects.filter(
                delete_time=CommonConstant.DELETE_TIME_ZERO,
                archive_time__gt=CommonConstant.DELETE_TIME_ZERO
            ),
            {'archive_date': 'archive_time'}
        )

        paginator = Paginator(queryset.order_by('-archive_time'), limit)
        page_obj = paginator.get_page(page)

        data_list = []
        for contract in page_obj:
            data_list.append({
                'id': contract.id,
                'code': contract.code or '',
                'name': contract.name or '',
                'customer': contract.customer or '',
                'cost': str(contract.cost or 0),
                'archive_time': timestamp_to_date(contract.archive_time),
                'archive_uid': contract.archive_uid or '',
                'create_time': _format_datetime_value(contract.create_time)
            })

        return JsonResponse({
            'code': ApiResponseCode.CODE_SUCCESS,
            'msg': 'success',
            'count': paginator.count,
            'data': data_list
        })

    def post(self, request):
        return self.get_data_list(request)


class StopListView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        if _is_data_request(request):
            return self.get_data_list(request)
        return render(request, 'contract/terminate_list.html')

    def get_data_list(self, request):
        page = safe_int(request.GET.get('page'), 1)
        limit = safe_int(
            request.GET.get('limit'),
            CommonConstant.DEFAULT_PAGE_SIZE)

        queryset = _filter_contract_queryset(
            request,
            Contract.objects.filter(
                delete_time=CommonConstant.DELETE_TIME_ZERO,
                stop_time__gt=CommonConstant.DELETE_TIME_ZERO
            ),
            {'terminate_date': 'stop_time'}
        )

        paginator = Paginator(queryset.order_by('-stop_time'), limit)
        page_obj = paginator.get_page(page)

        data_list = []
        for contract in page_obj:
            data_list.append({
                'id': contract.id,
                'code': contract.code or '',
                'name': contract.name or '',
                'customer': contract.customer or '',
                'cost': str(contract.cost or 0),
                'stop_time': timestamp_to_date(contract.stop_time),
                'stop_uid': contract.stop_uid or '',
                'stop_remark': contract.stop_remark or '',
                'create_time': _format_datetime_value(contract.create_time)
            })

        return JsonResponse({
            'code': ApiResponseCode.CODE_SUCCESS,
            'msg': 'success',
            'count': paginator.count,
            'data': data_list
        })

    def post(self, request):
        return self.get_data_list(request)


class VoidListView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        if _is_data_request(request):
            return self.get_data_list(request)
        return render(request, 'contract/cancel_list.html')

    def get_data_list(self, request):
        page = safe_int(request.GET.get('page'), 1)
        limit = safe_int(
            request.GET.get('limit'),
            CommonConstant.DEFAULT_PAGE_SIZE)

        queryset = _filter_contract_queryset(
            request,
            Contract.objects.filter(
                delete_time=CommonConstant.DELETE_TIME_ZERO,
                void_time__gt=CommonConstant.DELETE_TIME_ZERO
            ),
            {'cancel_date': 'void_time'}
        )

        paginator = Paginator(queryset.order_by('-void_time'), limit)
        page_obj = paginator.get_page(page)

        data_list = []
        for contract in page_obj:
            data_list.append({
                'id': contract.id,
                'code': contract.code or '',
                'name': contract.name or '',
                'customer': contract.customer or '',
                'cost': str(contract.cost or 0),
                'void_time': timestamp_to_date(contract.void_time),
                'void_uid': contract.void_uid or '',
                'void_remark': contract.void_remark or '',
                'create_time': _format_datetime_value(contract.create_time)
            })

        return JsonResponse({
            'code': ApiResponseCode.CODE_SUCCESS,
            'msg': 'success',
            'count': paginator.count,
            'data': data_list
        })

    def post(self, request):
        return self.get_data_list(request)


class ContractView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        if _is_data_request(request):
            return self.get_data_list(request)
        return render(request, 'contract/sales_list.html')

    def get_data_list(self, request):
        page = safe_int(request.GET.get('page'), 1)
        limit = safe_int(
            request.GET.get('limit'),
            CommonConstant.DEFAULT_PAGE_SIZE)

        queryset = _filter_contract_queryset(
            request,
            _base_active_contract_queryset()
        )

        paginator = Paginator(queryset.order_by('-create_time'), limit)
        page_obj = paginator.get_page(page)

        data_list = []
        for contract in page_obj:
            data_list.append({
                'id': contract.id,
                'code': contract.code or '',
                'name': contract.name or '',
                'customer': contract.customer or '',
                'cost': str(contract.cost or 0),
                'check_status': contract.check_status,
                'create_time': _format_datetime_value(contract.create_time),
                'sign_time': timestamp_to_date(contract.sign_time),
                'start_time': timestamp_to_date(contract.start_time),
                'end_time': timestamp_to_date(contract.end_time)
            })

        return JsonResponse({
            'code': ApiResponseCode.CODE_SUCCESS,
            'msg': 'success',
            'count': paginator.count,
            'data': data_list
        })

    def post(self, request):
        return self.get_data_list(request)


class ContractAddView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        return render(request, 'contract/sales/add.html', {
            'contract': None,
            'category_id': request.GET.get('category_id', '')
        })

    def post(self, request):
        try:
            payload = _build_contract_payload(request.POST.dict(), request.user)
            Contract.objects.create(**payload)
            logger.info(f"用户 {request.user.id} 添加了合同")
            return JsonResponse({
                'code': ApiResponseCode.CODE_SUCCESS,
                'msg': '添加成功'
            })
        except Exception as e:
            logger.error(f'添加合同失败: {str(e)}', exc_info=True)
            return JsonResponse({
                'code': ApiResponseCode.CODE_ERROR,
                'msg': '添加失败，请检查合同信息后重试'
            })


class ContractDetailView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request, id):
        contract = get_object_or_404(
            Contract,
            id=id,
            delete_time=CommonConstant.DELETE_TIME_ZERO
        )
        return render(request, 'contract/view.html', {'contract': contract})


class ContractUpdateView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def post(self, request, id):
        try:
            data = _parse_json_request(request)
            contract = get_object_or_404(
                Contract,
                id=id,
                delete_time=CommonConstant.DELETE_TIME_ZERO
            )
            payload = _build_contract_payload(data, request.user, require_all=False)

            if not payload:
                return JsonResponse({
                    'code': ApiResponseCode.CODE_ERROR,
                    'msg': '没有可更新的合同信息'
                })

            for key, value in payload.items():
                setattr(contract, key, value)

            contract.save()
            logger.info(f"用户 {request.user.id} 更新了合同: {id}")

            return JsonResponse({
                'code': ApiResponseCode.CODE_SUCCESS,
                'msg': '更新成功'
            })
        except Exception as e:
            logger.error(f'更新合同失败: {str(e)}', exc_info=True)
            return JsonResponse({
                'code': ApiResponseCode.CODE_ERROR,
                'msg': '更新失败，请检查合同信息后重试'
            })


class ContractDeleteView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def post(self, request):
        try:
            data = _parse_json_request(request)
            ids = _parse_int_ids(data.get('ids') or data.get('id'))

            if not ids:
                return JsonResponse({
                    'code': ApiResponseCode.CODE_ERROR,
                    'msg': '缺少合同ID'
                })

            current_time = int(time.time())
            updated_count = Contract.objects.filter(
                id__in=ids,
                delete_time=CommonConstant.DELETE_TIME_ZERO
            ).update(delete_time=current_time)

            if updated_count == 0:
                return JsonResponse({
                    'code': ApiResponseCode.CODE_ERROR,
                    'msg': '未找到可删除的合同记录'
                })

            logger.info(f"用户 {request.user.id} 删除了合同: {ids}")

            return JsonResponse({
                'code': ApiResponseCode.CODE_SUCCESS,
                'msg': '删除成功'
            })
        except Exception as e:
            logger.error(f'删除合同失败: {str(e)}', exc_info=True)
            return JsonResponse({
                'code': ApiResponseCode.CODE_ERROR,
                'msg': '删除失败，请确认合同状态后重试'
            })


class ContractExportView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request, export_type='sales'):
        if export_type == 'sales':
            queryset = _base_active_contract_queryset()
        else:
            queryset = Contract.objects.filter(delete_time=CommonConstant.DELETE_TIME_ZERO)
        title = '销售合同'
        filename_prefix = '销售合同'
        extra_columns = []
        date_field_map = None

        if export_type == 'archive':
            queryset = queryset.filter(archive_time__gt=CommonConstant.DELETE_TIME_ZERO)
            title = '归档合同'
            filename_prefix = '归档合同'
            date_field_map = {'archive_date': 'archive_time'}
            extra_columns = [
                ('归档时间', lambda contract: timestamp_to_date(contract.archive_time)),
                ('归档人ID', lambda contract: contract.archive_uid or '')
            ]
        elif export_type == 'terminate':
            queryset = queryset.filter(stop_time__gt=CommonConstant.DELETE_TIME_ZERO)
            title = '终止合同'
            filename_prefix = '终止合同'
            date_field_map = {'terminate_date': 'stop_time'}
            extra_columns = [
                ('终止时间', lambda contract: timestamp_to_date(contract.stop_time)),
                ('终止人ID', lambda contract: contract.stop_uid or ''),
                ('终止原因', lambda contract: contract.stop_remark or '')
            ]
        elif export_type == 'cancel':
            queryset = queryset.filter(void_time__gt=CommonConstant.DELETE_TIME_ZERO)
            title = '作废合同'
            filename_prefix = '作废合同'
            date_field_map = {'cancel_date': 'void_time'}
            extra_columns = [
                ('作废时间', lambda contract: timestamp_to_date(contract.void_time)),
                ('作废人ID', lambda contract: contract.void_uid or ''),
                ('作废原因', lambda contract: contract.void_remark or '')
            ]

        queryset = _filter_contract_queryset(request, queryset, date_field_map).order_by('-create_time')
        headers, rows = _build_contract_export_rows(queryset, extra_columns)
        return _build_xlsx_response(title, headers, rows, filename_prefix)


class PurchaseExportView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        queryset = _filter_purchase_queryset(
            request,
            Purchase.objects.select_related('cate').filter(delete_time__isnull=True)
        ).order_by('-create_time')
        headers = ['ID', '采购编号', '采购名称', '采购分类', '采购金额', '审核状态', '签订日期', '开始日期', '结束日期', '创建时间']
        status_map = {0: '待审核', 1: '审核中', 2: '审核通过', 3: '审核不通过', 4: '撤销审核'}
        rows = []
        for purchase in queryset:
            rows.append([
                purchase.id,
                purchase.code or '',
                purchase.name or '',
                purchase.cate.title if purchase.cate else '',
                float(_to_decimal(purchase.amount)),
                status_map.get(purchase.check_status, purchase.check_status),
                purchase.sign_time.strftime('%Y-%m-%d') if purchase.sign_time else '',
                purchase.start_time.strftime('%Y-%m-%d') if purchase.start_time else '',
                purchase.end_time.strftime('%Y-%m-%d') if purchase.end_time else '',
                _format_datetime_value(purchase.create_time),
            ])
        return _build_xlsx_response('采购合同', headers, rows, '采购合同')


class ContractArchiveView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def post(self, request, action=None):
        try:
            data, contract = _parse_contract_operation_request(request, action)
            action = data.get('action', 'archive')

            if action not in ['archive', 'unarchive']:
                return JsonResponse({
                    'code': ApiResponseCode.CODE_ERROR,
                    'msg': '无效的归档操作'
                })

            current_time = int(time.time())

            if action == 'archive':
                contract.archive_time = current_time
                contract.archive_uid = request.user.id
                msg = '归档成功'
            else:
                contract.archive_time = 0
                contract.archive_uid = 0
                msg = '取消归档成功'

            contract.save()
            logger.info(
                f"用户 {request.user.id} 对合同 {contract.id} 执行了{action}操作")

            return JsonResponse({
                'code': ApiResponseCode.CODE_SUCCESS,
                'msg': msg
            })
        except Exception as e:
            logger.error(f'归档操作失败: {str(e)}', exc_info=True)
            return JsonResponse({
                'code': ApiResponseCode.CODE_ERROR,
                'msg': '归档操作失败，请确认合同状态后重试'
            })


class ContractTerminateView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def post(self, request):
        try:
            data, contract = _parse_contract_operation_request(request)
            remark = data.get('remark', '')

            if contract.stop_time and contract.stop_time > CommonConstant.DELETE_TIME_ZERO:
                return JsonResponse({
                    'code': ApiResponseCode.CODE_ERROR,
                    'msg': '合同已终止'
                })

            contract.stop_time = int(time.time())
            contract.stop_uid = request.user.id
            contract.stop_remark = remark
            contract.save()

            logger.info(f"用户 {request.user.id} 终止了合同: {contract.id}")

            return JsonResponse({
                'code': ApiResponseCode.CODE_SUCCESS,
                'msg': '终止成功'
            })
        except Exception as e:
            logger.error(f'终止合同失败: {str(e)}', exc_info=True)
            return JsonResponse({
                'code': ApiResponseCode.CODE_ERROR,
                'msg': '终止操作失败，请确认合同状态后重试'
            })


class ContractCancelView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def post(self, request):
        try:
            data, contract = _parse_contract_operation_request(request)
            remark = data.get('remark', '')

            if contract.void_time and contract.void_time > CommonConstant.DELETE_TIME_ZERO:
                return JsonResponse({
                    'code': ApiResponseCode.CODE_ERROR,
                    'msg': '合同已作废'
                })

            contract.void_time = int(time.time())
            contract.void_uid = request.user.id
            contract.void_remark = remark
            contract.save()

            logger.info(f"用户 {request.user.id} 作废了合同: {contract.id}")

            return JsonResponse({
                'code': ApiResponseCode.CODE_SUCCESS,
                'msg': '作废成功'
            })
        except Exception as e:
            logger.error(f'作废合同失败: {str(e)}', exc_info=True)
            return JsonResponse({
                'code': ApiResponseCode.CODE_ERROR,
                'msg': '作废操作失败，请确认合同状态后重试'
            })




class ContractCategoryChildrenView(LoginRequiredMixin, View):
    login_url = '/user/login/'

    def get(self, request, parent_id=None):
        try:
            if parent_id is None or parent_id == '0':
                categories = ContractCategory.objects.filter(
                    parent__isnull=True,
                    is_active=True
                ).order_by('sort_order', 'id')
            else:
                categories = ContractCategory.objects.filter(
                    parent_id=parent_id,
                    is_active=True
                ).order_by('sort_order', 'id')

            data = []
            for category in categories:
                has_children = ContractCategory.objects.filter(
                    parent_id=category.id,
                    is_active=True
                ).exists()

                data.append({
                    'id': category.id,
                    'name': category.name,
                    'code': category.code,
                    'has_children': has_children
                })

            return JsonResponse({
                'code': 0,
                'msg': '',
                'data': data
            })

        except Exception as e:
            logger.error(f'获取合同分类失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '获取分类失败，请刷新页面后重试'})


class ContractCategoryView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        if _is_data_request(request):
            return self.get_data_list(request)
        return render(request, 'contract/contract_category_list.html')

    def get_data_list(self, request):
        return _serialize_basic_data_response(
            _filter_active(ContractCategory.objects.all()),
            request,
            lambda category: {
                'id': category.id,
                'name': category.name,
                'code': category.code,
                'parent': category.parent.name if category.parent else '',
                'description': category.description or '',
                'template_path': category.template_path or '',
                'sort_order': category.sort_order or 0,
                'is_active': bool(category.is_active),
                'created_at': category.created_at.strftime('%Y-%m-%d %H:%M:%S') if category.created_at else ''
            }
        )


class ContractCategoryAddView(LoginRequiredMixin, View):
    login_url = '/user/login/'

    def get(self, request):
        form = ContractCategoryForm()
        return render(request, 'contract/contract_category_form.html', {'form': form, 'object': None})

    def post(self, request):
        form = ContractCategoryForm(request.POST)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})
        try:
            form.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'保存合同分类失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查合同分类信息后重试'})


class ContractCategoryEditView(LoginRequiredMixin, View):
    login_url = '/user/login/'

    def get(self, request, id):
        category = _get_active_object(ContractCategory, id)
        form = ContractCategoryForm(instance=category)
        return render(request, 'contract/contract_category_form.html', {'form': form, 'object': category})

    def post(self, request, id):
        category = _get_active_object(ContractCategory, id)
        form = ContractCategoryForm(request.POST, instance=category)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})
        try:
            form.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'保存合同分类失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查合同分类信息后重试'})


class ProductCategoryView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        if _is_data_request(request):
            return self.get_data_list(request)
        return render(request, 'contract/product_category_list.html')

    def get_data_list(self, request):
        return _serialize_basic_data_response(
            _filter_active(ProductCategory.objects.all()),
            request,
            lambda cate: {
                'id': cate.id,
                'name': cate.name,
                'code': cate.code,
                'parent': cate.parent.name if cate.parent else '',
                'description': cate.description or '',
                'sort_order': cate.sort_order or 0,
                'is_active': bool(cate.is_active),
                'created_at': cate.created_at.strftime('%Y-%m-%d %H:%M:%S') if cate.created_at else ''
            }
        )


class ProductCategoryAddView(LoginRequiredMixin, View):
    login_url = '/user/login/'

    def get(self, request):
        form = ProductCategoryForm()
        return render(request, 'contract/product_category_form.html', {'form': form, 'object': None})

    def post(self, request):
        form = ProductCategoryForm(request.POST)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})
        try:
            form.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'保存产品分类失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查产品分类信息后重试'})


class ProductCategoryEditView(LoginRequiredMixin, View):
    login_url = '/user/login/'

    def get(self, request, id):
        category = _get_active_object(ProductCategory, id)
        form = ProductCategoryForm(instance=category)
        return render(request, 'contract/product_category_form.html', {'form': form, 'object': category})

    def post(self, request, id):
        category = _get_active_object(ProductCategory, id)
        form = ProductCategoryForm(request.POST, instance=category)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})
        try:
            form.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'保存产品分类失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查产品分类信息后重试'})


class ServiceCategoryView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        if _is_data_request(request):
            return self.get_data_list(request)
        return render(request, 'contract/service_category_list.html')

    def get_data_list(self, request):
        return _serialize_basic_data_response(
            _filter_active(ServiceCategory.objects.all()),
            request,
            lambda cate: {
                'id': cate.id,
                'name': cate.name,
                'code': cate.code,
                'parent': cate.parent.name if cate.parent else '',
                'description': cate.description or '',
                'sort_order': cate.sort_order or 0,
                'is_active': bool(cate.is_active),
                'created_at': cate.created_at.strftime('%Y-%m-%d %H:%M:%S') if cate.created_at else ''
            }
        )


class ServiceCategoryAddView(LoginRequiredMixin, View):
    login_url = '/user/login/'

    def get(self, request):
        form = ServiceCategoryForm()
        return render(request, 'contract/service_category_form.html', {'form': form, 'object': None})

    def post(self, request):
        form = ServiceCategoryForm(request.POST)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})
        try:
            form.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'保存服务分类失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查服务分类信息后重试'})


class ServiceCategoryEditView(LoginRequiredMixin, View):
    login_url = '/user/login/'

    def get(self, request, id):
        category = _get_active_object(ServiceCategory, id)
        form = ServiceCategoryForm(instance=category)
        return render(request, 'contract/service_category_form.html', {'form': form, 'object': category})

    def post(self, request, id):
        category = _get_active_object(ServiceCategory, id)
        form = ServiceCategoryForm(request.POST, instance=category)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})
        try:
            form.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'保存服务分类失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查服务分类信息后重试'})


class SupplierView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        if _is_data_request(request):
            return self.get_data_list(request)
        return render(request, 'contract/supplier_list.html')

    def get_data_list(self, request):
        return _serialize_basic_data_response(
            _filter_active(Supplier.objects.all()),
            request,
            lambda supplier: {
                'id': supplier.id,
                'name': supplier.name or '',
                'code': supplier.code or '',
                'contact': supplier.contact_person or '',
                'contact_person': supplier.contact_person or '',
                'phone': supplier.contact_phone or '',
                'contact_phone': supplier.contact_phone or '',
                'email': supplier.contact_email or '',
                'address': supplier.address or '',
                'sort_order': 0,
                'is_active': bool(supplier.is_active),
                'created_at': supplier.created_at.strftime('%Y-%m-%d %H:%M:%S') if supplier.created_at else ''
            },
            search_fields=['name', 'code', 'contact_person']
        )


class SupplierAddView(LoginRequiredMixin, View):
    login_url = '/user/login/'

    def get(self, request):
        form = SupplierForm()
        return render(request, 'contract/supplier_form.html', {'form': form, 'object': None})

    def post(self, request):
        form = SupplierForm(request.POST)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})
        try:
            form.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'保存供应商失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查供应商信息后重试'})


class SupplierEditView(LoginRequiredMixin, View):
    login_url = '/user/login/'

    def get(self, request, id):
        supplier = _get_active_object(Supplier, id)
        form = SupplierForm(instance=supplier)
        return render(request, 'contract/supplier_form.html', {'form': form, 'object': supplier})

    def post(self, request, id):
        supplier = _get_active_object(Supplier, id)
        form = SupplierForm(request.POST, instance=supplier)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})
        try:
            form.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'保存供应商失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查供应商信息后重试'})


class PurchaseCategoryView(LoginRequiredMixin, View):
    login_url = '/user/login/'
    redirect_field_name = 'next'

    def get(self, request):
        if _is_data_request(request):
            return self.get_data_list(request)
        return render(request, 'contract/purchase_category_list.html')

    def get_data_list(self, request):
        return _serialize_basic_data_response(
            _filter_active(PurchaseCategory.objects.all()),
            request,
            lambda cate: {
                'id': cate.id,
                'name': cate.name or '',
                'code': cate.code or '',
                'parent': cate.parent.name if cate.parent else '',
                'description': cate.description or '',
                'sort_order': cate.sort_order or 0,
                'is_active': bool(cate.is_active),
                'created_at': cate.created_at.strftime('%Y-%m-%d %H:%M:%S') if cate.created_at else ''
            }
        )


class PurchaseCategoryAddView(LoginRequiredMixin, View):
    login_url = '/user/login/'

    def get(self, request):
        form = PurchaseCategoryForm()
        return render(request, 'contract/purchase_category_form.html', {'form': form, 'object': None})

    def post(self, request):
        form = PurchaseCategoryForm(request.POST)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})
        try:
            form.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'保存采购分类失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查采购分类信息后重试'})


class PurchaseCategoryEditView(LoginRequiredMixin, View):
    login_url = '/user/login/'

    def get(self, request, id):
        category = _get_active_object(PurchaseCategory, id)
        form = PurchaseCategoryForm(instance=category)
        return render(request, 'contract/purchase_category_form.html', {'form': form, 'object': category})

    def post(self, request, id):
        category = _get_active_object(PurchaseCategory, id)
        form = PurchaseCategoryForm(request.POST, instance=category)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})
        try:
            form.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'保存采购分类失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查采购分类信息后重试'})


class PurchaseItemView(LoginRequiredMixin, View):
    login_url = '/user/login/'

    def get(self, request):
        if _is_data_request(request):
            return self.get_data_list(request)
        return render(request, 'contract/purchase_item_list.html')

    def get_data_list(self, request):
        return _serialize_basic_data_response(
            _filter_active(PurchaseItem.objects.select_related('category', 'supplier')),
            request,
            lambda item: {
                'id': item.id,
                'name': item.name or '',
                'code': item.code or '',
                'category': item.category.name if item.category else '',
                'specification': item.specification or '',
                'description': item.description or '',
                'unit': item.unit or '',
                'reference_price': str(item.reference_price) if item.reference_price else '0',
                'supplier': item.supplier.name if item.supplier else '',
                'sort_order': 0,
                'is_active': bool(item.is_active),
                'created_at': item.created_at.strftime('%Y-%m-%d %H:%M:%S') if item.created_at else ''
            }
        )


class PurchaseItemAddView(LoginRequiredMixin, View):
    login_url = '/user/login/'

    def get(self, request):
        form = PurchaseItemForm()
        return render(request, 'contract/purchase_item_form.html', {'form': form, 'object': None})

    def post(self, request):
        form = PurchaseItemForm(request.POST)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})
        try:
            form.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'保存采购项目失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查采购项目信息后重试'})


class PurchaseItemEditView(LoginRequiredMixin, View):
    login_url = '/user/login/'

    def get(self, request, id):
        item = _get_active_object(PurchaseItem, id)
        form = PurchaseItemForm(instance=item)
        return render(request, 'contract/purchase_item_form.html', {'form': form, 'object': item})

    def post(self, request, id):
        item = _get_active_object(PurchaseItem, id)
        form = PurchaseItemForm(request.POST, instance=item)
        if not form.is_valid():
            return JsonResponse({'code': 1, 'msg': '表单验证失败', 'errors': form.errors})
        try:
            form.save()
            return JsonResponse({'code': 0, 'msg': '保存成功'})
        except Exception as e:
            logger.error(f'保存采购项目失败: {str(e)}', exc_info=True)
            return JsonResponse({'code': 1, 'msg': '保存失败，请检查采购项目信息后重试'})


class ServiceListView(ServicesView):
    pass


class ServiceListAddView(ServicesAddView):
    pass


class ServiceListEditView(ServicesDetailView):
    pass
